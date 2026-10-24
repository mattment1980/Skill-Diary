"""Erzeugt template/Auswertung-Vorlage.pptx aus der eigenen PowerPoint-Vorlage (Entwicklungsgespräch).

Aufruf:  python3 tools/build_report_template.py <eigene-vorlage.pptx> template/Auswertung-Vorlage.pptx
Benötigt: pip install python-pptx pillow

Was passiert:
- Seite 1: Beispieltexte werden durch Platzhalter ersetzt ({{ZEITRAUM}}, {{SPIELER}}, {{B1}}…{{B4}},
  {{EINSCHAETZUNG}}, {{SUPER}}, {{VERBESSERN}}, {{STOLZ}}); Sterne werden als STAR_<Bereich>_<1..5> neu angelegt.
- Seite 2 (neu): gleicher Stil mit Kästen und Grafikflächen CHART_OVERVIEW/RADAR/TIMELINE/SKILLS.
- Bilder werden verkleinert, das Vorschaubild geleert (enthielt Beispieltexte).

Die Formnamen ("Textfeld 12", "Grafik 76" …) beziehen sich auf die aktuelle Vorlage. Ändert sich die
Vorlage, müssen diese Namen hier angepasst werden. js/export/report.js erwartet: Stern voll = rId3, leer = rId4.
WICHTIG: Die Quell-Vorlage enthält echte Beispieltexte – nicht ins Repo legen, nur das Ergebnis.
"""
import copy, io, sys, zipfile
from pptx import Presentation
from pptx.util import Cm, Emu
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
s1 = p.slides[0]
by = {sh.name: sh for sh in s1.shapes}

def set_text(shape, text):
    """Ganzen Text auf den ersten Run setzen (Formatierung bleibt), weitere Absätze/Runs entfernen."""
    tf = shape.text_frame
    paras = tf.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    runs = paras[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

# Seite 1: Platzhalter
by['Textfeld 11']._element.getparent().remove(by['Textfeld 11']._element)  # verdeckter Altrest
set_text(by['Textfeld 12'], 'Zeitraum:\t\t{{ZEITRAUM}}')
set_text(by['Textfeld 2'], 'Spieler:\t{{SPIELER}}')
for name, ph in (('Textfeld 4', 'B1'), ('Textfeld 6', 'B2'), ('Textfeld 8', 'B3'), ('Textfeld 10', 'B4')):
    set_text(by[name], '{{%s}}:' % ph)
set_text(by['Textfeld 53'], '{{EINSCHAETZUNG}}')
set_text(by['Textfeld 59'], '{{SUPER}}')
set_text(by['Textfeld 64'], '{{VERBESSERN}}')
set_text(by['Textfeld 69'], '{{STOLZ}}')

# Sterne neu: je Bereich 5 Bilder STAR_<bereich>_<n> (voll = Bild 2, leer = Bild 3)
star_full = by['Grafik 76']   # Vorlage für Bildelement
full_rid = star_full._element.blipFill.blip.rEmbed if hasattr(star_full._element, 'blipFill') else None
for n in ('Gruppieren 29', 'Gruppieren 30', 'Gruppieren 36', 'Gruppieren 42', 'Grafik 3', 'Grafik 7', 'Grafik 14', 'Grafik 78'):
    by[n]._element.getparent().remove(by[n]._element)
STEP = 281759
rows = {1: (2074563, 2679759), 2: (2074563, 2989921), 3: (4715886, 2679759), 4: (4715886, 2989921)}
tree = s1.shapes._spTree
next_id = 900
for b, (x0, y0) in rows.items():
    for i in range(5):
        el = copy.deepcopy(star_full._element)
        el.nvPicPr.cNvPr.set('name', f'STAR_{b}_{i+1}')
        el.nvPicPr.cNvPr.set('id', str(next_id)); next_id += 1
        for ext in el.nvPicPr.cNvPr.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}extLst'):
            el.nvPicPr.cNvPr.remove(ext)
        el.spPr.xfrm.off.x = x0 + i * STEP
        el.spPr.xfrm.off.y = y0
        tree.append(el)
star_full._element.getparent().remove(star_full._element)

# Seite 2: Auswertung im selben Stil
s2 = p.slides.add_slide(p.slide_layouts[6])
for ph in list(s2.placeholders):
    ph._element.getparent().remove(ph._element)
hdr = by['Grafik 16']
img_part = s1.part.related_part(hdr._element.blipFill.blip.rEmbed)
rid = s2.part.relate_to(img_part, RT.IMAGE)
t2 = s2.shapes._spTree
def clone(shape, name=None, x=None, y=None, w=None, h=None, text=None):
    el = copy.deepcopy(shape._element)
    t2.append(el)
    sh = s2.shapes[-1]
    if name: sh.name = name
    if x is not None: sh.left = Cm(x)
    if y is not None: sh.top = Cm(y)
    if w is not None: sh.width = Cm(w)
    if h is not None: sh.height = Cm(h)
    if text is not None: set_text(sh, text)
    return sh
pic = clone(hdr); pic._element.blipFill.blip.rEmbed = rid
clone(by['Textfeld 12']); clone(by['Textfeld 2'])
frame, title = by['Rechteck: abgerundete Ecken 13'], by['Textfeld 15']
boxes = [  # Titel, Rahmen x,y,w,h, Bildname, Bild x,y,w,h
    ('Überblick:',          0.9, 6.4, 8.4, 7.0, 'CHART_OVERVIEW', 1.6, 7.1, 7.0, 5.47),
    ('Niveau je Bereich:',  9.7, 6.4, 8.5, 7.0, 'CHART_RADAR',    10.4, 7.1, 7.1, 5.48),
    ('Verlauf:',            0.9, 14.2, 17.3, 5.2, 'CHART_TIMELINE', 1.3, 14.8, 16.5, 4.3),
    ('Niveau je Skill:',    0.9, 20.2, 17.3, 6.9, 'CHART_SKILLS',  1.95, 20.85, 15.2, 5.53),
]
for ttl, fx, fy, fw, fh, img, ix, iy, iw, ih in boxes:
    clone(frame, x=fx, y=fy, w=fw, h=fh)
    buf = io.BytesIO(); Image.new('RGB', (int(iw * 40), int(ih * 40)), 'white').save(buf, 'PNG'); buf.seek(0)
    ph = s2.shapes.add_picture(buf, Cm(ix), Cm(iy), Cm(iw), Cm(ih)); ph.name = img
    clone(title, x=fx + 1.2, y=fy - 0.4, text=ttl)
p.save(OUT)

# Bilder verkleinern, Vorschaubild ersetzen (enthielt Beispieltexte)
src = zipfile.ZipFile(OUT); data = {n: src.read(n) for n in src.namelist()}; src.close()
def shrink(name, width, colors=None):
    im = Image.open(io.BytesIO(data[name])); h = int(im.height * width / im.width)
    im = im.resize((width, h), Image.LANCZOS)
    if colors: im = im.convert('RGB').quantize(colors)
    b = io.BytesIO(); im.save(b, 'PNG', optimize=True); data[name] = b.getvalue()
shrink('ppt/media/image1.png', 1400, 256)
shrink('ppt/media/image2.png', 150); shrink('ppt/media/image3.png', 150)
b = io.BytesIO(); Image.new('RGB', (180, 256), 'white').save(b, 'JPEG'); data['docProps/thumbnail.jpeg'] = b.getvalue()
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as z:
    for n, d in data.items(): z.writestr(n, d)
