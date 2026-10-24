"""Erzeugt die leere Datenbank-Vorlage template/Skill-Diary-Vorlage.xlsx.

Aufruf:  python3 tools/build_xlsx_template.py template/Skill-Diary-Vorlage.xlsx
Benötigt: pip install openpyxl

Die App lädt diese Vorlage ins OneDrive, falls dort noch keine Datenbank existiert.
Spalten müssen zu CONFIG.feedbackColumns in js/core/config.js passen. Neue Spalten in bestehenden
Datenbanken ergänzt die App selbst (js/services/onedrive.js → migrateFeedbackColumns).
"""
import sys
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment

FB_COLS = ['ID', 'Zeitstempel', 'Datum', 'Art', 'Anlass', 'Spieler', 'Position', 'Bereich', 'Skill', 'Bewertung',
           'Kommentar', 'Quelle', 'Gruppe', 'VideoLink', 'Niveau']
KATALOG = {
    'Spielintelligenz': ['Scanning / Orientierung', 'Entscheidungsqualität', 'Raumnutzung', 'Verhalten ohne Ball'],
    'Technik': ['Ballkontrolle / 1st touch', 'Dribbling / 1vs1 offensiv', 'Zweikampf defensiv', 'Passspiel', 'Torschuss', 'Beidfüßigkeit'],
    'Koordination / Physis': ['Körperbeherrschung', 'Laufkoordination', 'Beweglichkeit / Gewandtheit', 'Rhythmus- / Kopplungsfähigkeit'],
    'Soziales / Mentales': ['Lernbereitschaft', 'Frustrationstoleranz', 'Kommunikation', 'Einsatzbereitschaft', 'Teamgeist'],
}
POS = [('Torwart', 'TW', 'Torwart'), ('Verteidigung', 'LV', 'Linker Verteidiger'), ('Verteidigung', 'RV', 'Rechter Verteidiger'),
       ('Mittelfeld', 'LM', 'Linkes Mittelfeld'), ('Mittelfeld', 'ZM', 'Zentrales Mittelfeld'), ('Mittelfeld', 'RM', 'Rechtes Mittelfeld'),
       ('Sturm', 'ST', 'Stürmer')]
HEAD, FILL = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='01663F')


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Feedback'
    ws.append(FB_COLS)
    ws.append([''] * len(FB_COLS))  # Excel-Tabellen brauchen mindestens eine Datenzeile
    last = chr(64 + len(FB_COLS))
    t = Table(displayName='Feedback', ref=f'A1:{last}2')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium7', showRowStripes=True)
    ws.add_table(t)
    for i, w in enumerate([10, 17, 11, 9, 16, 22, 9, 20, 26, 10, 70, 10, 12, 20, 8]):
        ws.column_dimensions[chr(65 + i)].width = w
    # Datum als Datum, Textspalten als Text (verhindert z. B. "1:1" -> Uhrzeit)
    for r in range(2, 3001):
        for col in 'ADEFGHIJKLMN':
            ws[f'{col}{r}'].number_format = '@'
        ws[f'B{r}'].number_format = 'DD.MM.YYYY HH:MM'
        ws[f'C{r}'].number_format = 'DD.MM.YYYY'
        ws[f'K{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    sheet(wb.create_sheet('Spieler'), ['Name', 'Aliase', 'Aktiv'], [], [26, 30, 8])
    sheet(wb.create_sheet('Kompetenzen'), ['Bereich', 'Skill'], [[b, s] for b, l in KATALOG.items() for s in l], [24, 32])
    sheet(wb.create_sheet('Positionen'), ['Gruppe', 'Position', 'Bezeichnung'], [list(p) for p in POS], [16, 10, 24])
    wb.save(path)


if __name__ == '__main__':
    build(sys.argv[1] if len(sys.argv) > 1 else 'template/Skill-Diary-Vorlage.xlsx')
